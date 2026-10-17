import json
import time

import pytest
from fastapi.testclient import TestClient
from google.genai import errors as genai_errors

from deed_extractor import api
from deed_extractor.evaluate import EvaluationReport
from deed_extractor.schema import ExtractionResult

KEY = {"X-Gemini-Key": "test-key"}


@pytest.fixture
def client():
    return TestClient(api.app)


def test_config(client):
    body = client.get("/api/config").json()
    assert body["models"] == api.GeminiExtractor.MODELS
    assert body["sample_available"] is True


def finish(client, response) -> dict:
    """Wait for a job started by response to end, and return its final progress."""
    assert response.status_code == 202, response.text
    job = response.json()
    for _ in range(250):
        job = client.get(f"/api/jobs/{job['id']}").json()
        if job["status"] == "done":
            return job
        time.sleep(0.02)
    raise AssertionError("job did not finish")


def result_of(client, job: dict, i: int = 0) -> dict:
    return client.get(f"/api/jobs/{job['id']}/documents/{job['documents'][i]['id']}").json()


def extraction(deed: dict, issues: list | None = None) -> ExtractionResult:
    return ExtractionResult(deed=deed, meta={"model": "m"}, validation=issues or [])


def test_jobs_require_a_key(client):
    for task in ("extract", "evaluate", "recheck"):
        assert client.post("/api/jobs", data={"use_sample": "true", "task": task}).status_code == 401


def test_jobs_reject_non_pdf_and_unknown_model(client):
    response = client.post("/api/jobs", headers=KEY, files=[("files", ("a.pdf", b"hello", "application/pdf"))])
    assert response.status_code == 415
    response = client.post("/api/jobs", headers=KEY, data={"use_sample": "true", "model": "gpt-4"})
    assert response.status_code == 400
    assert client.post("/api/jobs", headers=KEY, data={"task": "summarise"}).status_code == 422


def test_extract_sample(client, monkeypatch, sample_truth):
    seen = {}

    def fake_pipeline(pdf_path, api_key, model_name, on_stage=None):
        seen.update(api_key=api_key, model=model_name)
        on_stage("extracting")
        return extraction(sample_truth)

    monkeypatch.setattr(api, "run_pipeline", fake_pipeline)
    job = finish(client, client.post("/api/jobs", headers=KEY, data={"use_sample": "true"}))
    assert job["kind"] == "extract" and job["documents"][0]["status"] == "done"
    assert result_of(client, job)["deed"] == sample_truth
    assert seen == {"api_key": "test-key", "model": api.GeminiExtractor.MODELS[0]}


def fake_evaluation(seen: dict):
    def run(pdf_path, api_key, truth=None, on_stage=None, raise_if_all_failed=False):
        seen.update(truth=truth, raises=raise_if_all_failed)
        return EvaluationReport(pdf=pdf_path, models=[], raw_results={})
    return run


def test_evaluate_sample_uses_ground_truth(client, monkeypatch, sample_truth):
    seen = {}
    monkeypatch.setattr(api, "run_evaluation", fake_evaluation(seen))
    job = finish(client, client.post("/api/jobs", headers=KEY, data={"use_sample": "true", "task": "evaluate"}))
    assert seen == {"truth": sample_truth, "raises": True}
    assert job["kind"] == "evaluate" and "history_id" not in job["documents"][0]  # evaluations are not saved
    assert result_of(client, job)["pdf"] == "sample_deed.pdf"


def test_evaluate_upload_has_no_ground_truth(client, monkeypatch):
    seen = {}
    monkeypatch.setattr(api, "run_evaluation", fake_evaluation(seen))
    pdf = api.SAMPLE_PDF.read_bytes()
    files = [("files", ("mine.pdf", pdf, "application/pdf"))]
    job = finish(client, client.post("/api/jobs", headers=KEY, data={"task": "evaluate"}, files=files))
    assert seen["truth"] is None
    assert result_of(client, job)["pdf"] == "mine.pdf"
    two = files + [("files", ("other.pdf", pdf, "application/pdf"))]
    assert client.post("/api/jobs", headers=KEY, data={"task": "evaluate"}, files=two).status_code == 413


def test_evaluate_with_a_rejected_key_fails_the_job(client, monkeypatch):
    def rejected(**kwargs):
        raise genai_errors.APIError(400, {"error": {"message": "API key not valid. Please pass a valid API key."}})

    monkeypatch.setattr(api, "run_evaluation", rejected)
    job = finish(client, client.post("/api/jobs", headers=KEY, data={"use_sample": "true", "task": "evaluate"}))
    assert job["documents"][0]["error"] == "Gemini rejected the API key."


def test_recheck_validates_its_input(client, sample_truth):
    def post(**data):
        return client.post("/api/jobs", headers=KEY, data={"use_sample": "true", "task": "recheck", **data})

    assert post(deed="not json", fields="[]").status_code == 422
    assert post(deed=json.dumps(sample_truth), fields="[]").status_code == 422
    assert post(deed=json.dumps(sample_truth), fields=json.dumps(["property.colour"])).status_code == 422


def test_recheck_returns_answers(client, monkeypatch, sample_truth):
    def fake(pdf_path, api_key, deed, fields, model_name, on_stage=None):
        return {"answers": [{"field": fields[0], "value": "x", "agrees": False}]}

    monkeypatch.setattr(api, "run_recheck", fake)
    job = finish(client, client.post("/api/jobs", headers=KEY, data={
        "use_sample": "true", "task": "recheck",
        "deed": json.dumps(sample_truth), "fields": json.dumps(["property.kaek"]),
    }))
    assert job["kind"] == "recheck"
    assert result_of(client, job)["answers"][0]["field"] == "property.kaek"


def test_demo_deeds_need_no_key(client):
    demos = client.get("/api/config").json()["demos"]
    assert demos and all({"id", "title", "description"} <= set(d) for d in demos)
    result = client.get(f"/api/demo/{demos[0]['id']}").json()
    assert result["deed"]["sellers"] and result["meta"]["model"] and result["meta"]["note"]
    pdf = client.get(f"/api/demo/{demos[0]['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.get("/api/demo/nope").status_code == 404
    assert client.get("/api/demo/..%2Fsample_deed/pdf").status_code == 404


def test_demo_deeds_can_be_evaluated_with_their_answers(client, monkeypatch):
    seen = {}
    monkeypatch.setattr(api, "run_evaluation", fake_evaluation(seen))
    demo_id = client.get("/api/config").json()["demos"][0]["id"]
    finish(client, client.post("/api/jobs", headers=KEY, data={"demo": demo_id, "task": "evaluate"}))
    assert seen["truth"]["sellers"]
    assert client.post("/api/jobs", headers=KEY, data={"demo": "nope"}).status_code == 404


def test_validate_edited_deed(client, sample_truth):
    issues = client.post("/api/validate", json=sample_truth).json()["issues"]
    assert any(i["field"] == "sellers[0].tax_id" for i in issues)
    assert client.post("/api/validate", json={"sellers": "not a list"}).status_code == 422


def test_export_xlsx_and_csv(client, sample_truth):
    import io

    from openpyxl import load_workbook

    response = client.post("/api/export?format=xlsx", json=sample_truth)
    assert response.status_code == 200
    assert 'filename="deed.xlsx"' in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Deed", "People"]
    people = list(wb["People"].values)
    assert people[1][:3] == ("Seller", sample_truth["sellers"][0]["full_name"], sample_truth["sellers"][0]["tax_id"])
    assert len(people) == 1 + 2 + 1 + 1  # header, 2 sellers, 1 buyer, 1 representative

    csv_text = client.post("/api/export?format=csv", json=sample_truth).content.decode("utf-8-sig")
    assert "Transaction,Price (EUR),185000" in csv_text
    assert client.post("/api/export?format=pdf", json=sample_truth).status_code == 422


def test_exports_never_contain_formulas(client, sample_truth):
    """A value from the document that looks like a formula is written as text (CSV / formula injection)."""
    import io

    from openpyxl import load_workbook

    deed = {**sample_truth, "notary": {**sample_truth["notary"], "name": '=HYPERLINK("http://x","click")'}}
    wb = load_workbook(io.BytesIO(client.post("/api/export?format=xlsx", json=deed).content))
    names = [row[2] for row in wb["Deed"].values if row[1] == "Name"]
    assert names == ["'=HYPERLINK(\"http://x\",\"click\")"]
    csv_text = client.post("/api/export?format=csv", json=deed).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and ",=HYPERLINK" not in csv_text and ',"=HYPERLINK' not in csv_text
    batch = [{"name": "@evil.pdf", "deed": deed}]
    rows = list(load_workbook(io.BytesIO(
        client.post("/api/export/batch?format=xlsx", json={"deeds": batch}).content))["Deeds"].values)
    assert rows[1][0] == "'@evil.pdf"


def test_batch_job_end_to_end(client, monkeypatch, sample_truth):
    def fake_pipeline(pdf_path, api_key, model_name, on_stage=None):
        on_stage("reading")
        return extraction(sample_truth)

    monkeypatch.setattr(api, "run_pipeline", fake_pipeline)
    pdf = api.SAMPLE_PDF.read_bytes()
    files = [("files", ("one.pdf", pdf, "application/pdf")), ("files", ("two.pdf", pdf, "application/pdf"))]
    job = finish(client, client.post("/api/jobs", headers=KEY, files=files))
    assert [d["filename"] for d in job["documents"]] == ["one.pdf", "two.pdf"]
    assert all(d["status"] == "done" for d in job["documents"])
    assert result_of(client, job)["deed"]["notary"] == sample_truth["notary"]
    assert client.get("/api/jobs/nope").status_code == 404


def test_batch_job_checks_every_file(client):
    pdf = api.SAMPLE_PDF.read_bytes()
    files = [("files", ("ok.pdf", pdf, "application/pdf")), ("files", ("bad.pdf", b"hello", "application/pdf"))]
    response = client.post("/api/jobs", headers=KEY, files=files)
    assert response.status_code == 415
    assert "bad.pdf" in response.json()["detail"]
    too_many = [("files", (f"{i}.pdf", pdf, "application/pdf")) for i in range(api.MAX_BATCH_FILES + 1)]
    assert client.post("/api/jobs", headers=KEY, files=too_many).status_code == 413


def test_batch_export(client, sample_truth):
    import io

    from openpyxl import load_workbook

    deeds = [{"name": "a.pdf", "deed": sample_truth}, {"name": "b.pdf", "deed": sample_truth}]
    wb = load_workbook(io.BytesIO(client.post("/api/export/batch?format=xlsx", json={"deeds": deeds}).content))
    assert wb.sheetnames == ["Deeds", "People"]
    rows = list(wb["Deeds"].values)
    assert rows[0][0] == "File" and [r[0] for r in rows[1:]] == ["a.pdf", "b.pdf"]
    assert len(list(wb["People"].values)) == 1 + 2 * 4  # 3 parties + 1 representative per deed
    csv = client.post("/api/export/batch?format=csv", json={"deeds": deeds}).content.decode("utf-8-sig")
    assert csv.splitlines()[1].startswith("a.pdf,")
    assert client.post("/api/export/batch", json={"deeds": [{"name": "x"}]}).status_code == 422


def test_finished_jobs_go_to_history_without_the_pdf(client, monkeypatch, sample_truth):
    issues = [{"field": "property.kaek", "message": "x", "severity": "warning"}]
    monkeypatch.setattr(api, "run_pipeline", lambda pdf_path, api_key, model_name, on_stage=None:
                        extraction(sample_truth, issues))
    job = finish(client, client.post("/api/jobs", headers=KEY, data={"use_sample": "true"}))
    [doc] = job["documents"]
    history_id = doc["history_id"]
    assert result_of(client, job)["meta"]["history_id"] == history_id  # so the UI can save edits

    record = client.get(f"/api/history/{history_id}").json()
    assert record["filename"] == "sample_deed.pdf" and record["result"]["deed"] == sample_truth
    assert "pdf" not in {k.lower() for k in record}  # results only
    corrected = {**sample_truth, "transaction": {**sample_truth["transaction"], "price": 1}}
    saved = client.put(f"/api/history/{history_id}", json={"edited": corrected, "reviewed": ["property.kaek"]})
    assert saved.status_code == 200
    item = next(i for i in client.get("/api/history").json() if i["id"] == history_id)
    assert item["edited"] and item["summary"]["price"] == 1 and item["summary"]["to_check"] == 0
    assert client.delete(f"/api/history/{history_id}").status_code == 204
    assert client.get(f"/api/history/{history_id}").status_code == 404
