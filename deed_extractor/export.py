"""
Export an extracted (and possibly hand-corrected) deed to CSV or Excel.

CSV is one flat table of field / value rows. Excel has two sheets: the deed
details, and one row per person with their role. Many deeds export as one
row per deed, plus (in Excel) everyone from every deed with their file name.
"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DETAIL_FIELDS = [
    ("Notary", "notary", "name", "Name"),
    ("Notary", "notary", "location", "Location"),
    ("Notary", "notary", "registry_number", "Registry number"),
    ("Property", "property", "type", "Type"),
    ("Property", "property", "address", "Address"),
    ("Property", "property", "municipal_unit", "Municipal unit"),
    ("Property", "property", "floor", "Floor"),
    ("Property", "property", "area_sqm", "Area (m²)"),
    ("Property", "property", "kaek", "ΚΑΕΚ"),
    ("Property", "property", "building_permit", "Building permit"),
    ("Transaction", "transaction", "deed_type", "Deed type"),
    ("Transaction", "transaction", "price", "Price (EUR)"),
    ("Transaction", "transaction", "date", "Date"),
]
ROLES = [("sellers", "Seller"), ("buyers", "Buyer"), ("legal_representatives", "Representative")]
PEOPLE_COLUMNS = ["Role", "Name", "Tax ID (ΑΦΜ)", "ID document", "Address", "Represents"]

# Excel runs a cell starting with one of these as a formula (CSV / formula injection).
# Every value comes from the document or the user, so none is trusted: a leading
# apostrophe makes it plain text.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def cells(row: list) -> list:
    return ["'" + v if isinstance(v, str) and v.startswith(FORMULA_START) else v for v in row]


def detail_rows(deed: dict) -> list[list]:
    return [
        [section, label, (deed.get(key) or {}).get(field)]
        for section, key, field, label in DETAIL_FIELDS
    ]


def people_rows(deed: dict) -> list[list]:
    return [
        [role, p.get("full_name"), p.get("tax_id"), p.get("id_document"), p.get("address"), p.get("represents")]
        for key, role in ROLES
        for p in deed.get(key) or []
    ]


def to_csv(deed: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Section", "Field", "Value"])
    writer.writerows(cells(row) for row in detail_rows(deed))
    for key, role in ROLES:
        for i, p in enumerate(deed.get(key) or [], start=1):
            for field, value in p.items():
                writer.writerow(cells([f"{role} {i}", field, value]))
    # BOM so Excel opens the Greek text as UTF-8
    return ("﻿" + out.getvalue()).encode("utf-8")


def _style(sheet, widths: list[int]) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(i + 1)].width = width
    sheet.freeze_panes = "A2"


def _xlsx_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_xlsx(deed: dict) -> bytes:
    wb = Workbook()
    details = wb.active
    details.title = "Deed"
    details.append(["Section", "Field", "Value"])
    for row in detail_rows(deed):
        details.append(cells(row))

    people = wb.create_sheet("People")
    people.append(PEOPLE_COLUMNS)
    for row in people_rows(deed):
        people.append(cells(row))

    _style(details, [14, 18, 60])
    _style(people, [16, 40, 16, 16, 45, 40])
    return _xlsx_bytes(wb)

# ---------------------------------------------------------------------------
# Many deeds at once (batch or history): one row per deed
# ---------------------------------------------------------------------------

BATCH_COLUMNS = ["File", *(label for _, _, _, label in DETAIL_FIELDS), "Sellers", "Buyers"]


def batch_rows(deeds: list[tuple[str, dict]]) -> list[list]:
    def names(deed: dict, key: str) -> str:
        return "; ".join(p.get("full_name") or "?" for p in deed.get(key) or [])

    return [
        [name, *(row[2] for row in detail_rows(deed)), names(deed, "sellers"), names(deed, "buyers")]
        for name, deed in deeds
    ]


def batch_to_csv(deeds: list[tuple[str, dict]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(BATCH_COLUMNS)
    writer.writerows(cells(row) for row in batch_rows(deeds))
    return ("﻿" + out.getvalue()).encode("utf-8")


def batch_to_xlsx(deeds: list[tuple[str, dict]]) -> bytes:
    wb = Workbook()
    table = wb.active
    table.title = "Deeds"
    table.append(BATCH_COLUMNS)
    for row in batch_rows(deeds):
        table.append(cells(row))

    people = wb.create_sheet("People")
    people.append(["File", *PEOPLE_COLUMNS])
    for name, deed in deeds:
        for row in people_rows(deed):
            people.append(cells([name, *row]))

    _style(table, [28] + [18] * len(DETAIL_FIELDS) + [45, 45])
    _style(people, [28, 16, 40, 16, 16, 45, 40])
    return _xlsx_bytes(wb)
